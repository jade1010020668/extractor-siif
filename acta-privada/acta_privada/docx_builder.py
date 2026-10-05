"""ActaData -> .docx con el formato de actas de la Sala Plena.

Si existe plantilla/plantilla_acta.docx se reutilizan su encabezado, pie,
estilos y tamaño de página (se borra todo el contenido y los metadatos).
Sin plantilla se genera un documento equivalente desde cero.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from .models import ActaData, Person

PLANTILLA = Path(__file__).resolve().parent.parent / "plantilla" / "plantilla_acta.docx"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
FONT, SIZE = "Arial", 11
ANCHO = 9624  # twips, igual al acta de referencia


def fecha_larga(f, mayus=False) -> str:
    if not f:
        return "[FECHA]"
    s = f"{f[2]} de {MESES[f[1] - 1]} de {f[0]}"
    return s.upper() if mayus else s


# ─────────────────────────── helpers de formato ───────────────────────────
def _run(p, text, bold=False):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    r.font.size = Pt(SIZE)
    return r


def _para(cell, text="", bold=False, align=AL.JUSTIFY, reuse=False):
    """Añade un párrafo; con reuse=True usa el primero (vacío) de la celda."""
    p = cell.paragraphs[0] if reuse else cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    if text:
        _run(p, text, bold)
    return p


def _rich(cell, label, body, reuse=False):
    p = _para(cell, reuse=reuse)
    _run(p, label, True)
    _run(p, body)
    return p


def _blank(cell):
    _para(cell)


def _borders(tbl, outer=12, inner=4):
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge, sz in (("top", outer), ("left", outer), ("bottom", outer), ("right", outer),
                     ("insideH", inner), ("insideV", inner)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        b.append(e)
    tblPr.append(b)


def _table(doc, rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    _borders(t)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Twips(w)
    return t


def _cargo_firma(p: Person) -> str:
    """'Comisionada Presidente' -> 'Presidente' (como en el acta de referencia)."""
    return re.sub(r"^Comisionad[oa]\s+", "", p.titulo)


def _persona(cell, label, p: Person | None, reuse=False):
    _para(cell, label, True, AL.LEFT, reuse=reuse)
    _para(cell, p.nombre if p else "[COMPLETAR]", False, AL.LEFT)


def _seccion(cell, label, parrafos):
    """Primer párrafo con etiqueta en negrita, el resto sin ella."""
    for i, par in enumerate(parrafos):
        if i == 0:
            _rich(cell, label, par)
        else:
            _para(cell, par)
        _blank(cell)


# ───────────────────────────── construcción ─────────────────────────────
def _base_document():
    if PLANTILLA.exists():
        doc = Document(str(PLANTILLA))
        body = doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
    else:
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Twips(12240), Twips(18720)   # Oficio
        s.left_margin, s.right_margin = Twips(1418), Twips(1304)
        s.top_margin, s.bottom_margin = Twips(1700), Twips(709)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(SIZE)
    cp = doc.core_properties                      # sin autor ni metadatos heredados
    cp.author = cp.last_modified_by = cp.title = cp.subject = cp.comments = cp.keywords = ""
    return doc


def build_docx(acta: ActaData) -> bytes:
    doc = _base_document()

    for txt in (acta.tipo_sesion, f"ACTA No. {acta.numero or '___'}", fecha_larga(acta.fecha, True)):
        p = doc.add_paragraph()
        p.alignment = AL.CENTER
        p.paragraph_format.space_after = Pt(0)
        _run(p, txt, True)
    doc.add_paragraph()

    # Tabla 1: datos generales y asistentes
    t = _table(doc, 5, 2, (4605, 5019))
    _rich(t.cell(0, 0), "Ciudad: ", acta.ciudad, reuse=True)
    _rich(t.cell(1, 0), "Hora de Inicio: ", acta.hora_inicio or "[COMPLETAR]", reuse=True)
    _rich(t.cell(2, 0), "Hora de Finalización: ", acta.hora_fin or "[COMPLETAR]", reuse=True)
    rel = acta.relatoria
    _persona(t.cell(3, 0), f"{rel.titulo if rel else 'Profesional Universitaria'}:", rel, reuse=True)
    _para(t.cell(0, 1), "Comisionados", True, AL.CENTER, reuse=True)
    if acta.presidente:
        _persona(t.cell(1, 1), f"{acta.presidente.titulo}:", acta.presidente, reuse=True)
    else:
        _para(t.cell(1, 1), "[COMPLETAR PRESIDENTE]", reuse=True)
    for i, c in enumerate(acta.comisionados):
        if i < 2:
            _persona(t.cell(2 + i, 1), f"{c.titulo}:", c, reuse=True)
        else:                                       # más de 3 comisionados: se apilan
            _persona(t.cell(3, 1), f"{c.titulo}:", c)
    merged = t.cell(4, 0).merge(t.cell(4, 1))
    for extra in merged.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    _rich(merged, "Invitados: ", acta.invitados, reuse=True)
    doc.add_paragraph()

    # Tabla 2: orden del día
    t = _table(doc, 1, 1, (ANCHO,))
    c = t.cell(0, 0)
    _para(c, "ORDEN DEL DÍA", True, AL.CENTER, reuse=True)
    _blank(c)
    for i, item in enumerate(acta.orden_del_dia, 1):
        _para(c, f"{i}°. {item}", True)
        _blank(c)
    doc.add_paragraph()

    # Tabla 3: desarrollo (punto 1: aprobación del orden del día)
    t = _table(doc, 1, 1, (ANCHO,))
    c = t.cell(0, 0)
    _para(c, "DESARROLLO DEL ORDEN DEL DÍA", True, AL.CENTER, reuse=True)
    _blank(c)
    _para(c, f"1°. {acta.orden_del_dia[0] if acta.orden_del_dia else 'Aprobación del Orden del Día'}", True)
    _blank(c)
    for chunk in acta.apertura.split("\n\n"):
        if chunk.strip():
            _para(c, chunk.strip())
            _blank(c)
    doc.add_paragraph()

    # Puntos 2..n
    for punto in range(2, len(acta.orden_del_dia) + 1):
        t = _table(doc, 1, 1, (ANCHO,))
        c = t.cell(0, 0)
        _para(c, f"{punto}°. {acta.orden_del_dia[punto - 1]}", True, reuse=True)
        _blank(c)
        for k, tema in enumerate([x for x in acta.temas if x.punto == punto], 1):
            _para(c, f"{punto}.{k}. {tema.titulo}", True)
            _blank(c)
            if tema.pretension:
                _rich(c, "Pretensión: ", tema.pretension)
                _blank(c)
            _seccion(c, "Presentación del tema: ", tema.presentacion)
            _seccion(c, "Discusión sobre el tema: ", tema.discusion)
            _seccion(c, "Solicitud: ", tema.solicitud)
            if tema.decision:
                _rich(c, "Decisión: ", tema.decision)
                _blank(c)
            for par in tema.cierre:
                _para(c, par)
                _blank(c)
        doc.add_paragraph()

    # Firmas
    t = _table(doc, 2, 1, (ANCHO,))
    _para(t.cell(0, 0), "Firmas", True, AL.CENTER, reuse=True)
    c = t.cell(1, 0)
    _para(c, acta.nota_firmas, reuse=True)
    for _ in range(3):
        _blank(c)
    firmantes = [x for x in (acta.presidente, acta.relatoria) if x]
    _para(c, "\t\t".join(x.nombre.upper() for x in firmantes), True, AL.LEFT)
    _para(c, "\t\t".join(_cargo_firma(x) for x in firmantes), False, AL.LEFT)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
