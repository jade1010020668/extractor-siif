"""Lectura de transcripciones (Teams/Word u otras con el mismo patrón):

    Nombre del hablante   0:15
    Texto de la intervención...
"""
from __future__ import annotations

import io
import re

from .models import Segment, Transcript

MESES = {m: i for i, m in enumerate(
    ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
     "septiembre", "octubre", "noviembre", "diciembre"], 1)}

_SPEAKER = re.compile(r"^(?P<spk>\S.*?)(?:\s{2,}|\t)(?P<t>\d{1,2}:\d{2}(?::\d{2})?)\s*$")
_NOISE = re.compile(r"\bha (iniciado|detenido) la transcripci[oó]n\b", re.I)


def _to_seconds(t: str) -> int:
    parts = [int(x) for x in t.split(":")]
    while len(parts) < 3:
        parts.insert(0, 0)
    h, m, s = parts
    return h * 3600 + m * 60 + s


def _parse_duration(line: str) -> int | None:
    m = re.fullmatch(r"\s*(?:(\d+)\s*h\s*)?(?:(\d+)\s*min\s*)?(?:(\d+)\s*s)?\s*", line)
    if m and any(m.groups()):
        h, mi, s = (int(x or 0) for x in m.groups())
        return h * 3600 + mi * 60 + s
    return None


def _parse_title(title: str, tr: Transcript) -> None:
    m = re.search(r"(\d{1,2}) de (\w+) de (\d{4})", title, re.I)
    if m and m.group(2).lower() in MESES:
        tr.fecha = (int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1)))
    m = re.search(r"a las (\d{1,2})[:.]?(\d{2})\s*([ap])\.?\s*m", title, re.I)
    if m:
        h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).lower()
        h = h % 12 + (12 if ap == "p" else 0)
        tr.hora_inicio = (h, mi)


def parse_text(text: str) -> Transcript:
    tr = Transcript()
    header: list[str] = []
    cur: Segment | None = None
    for line in text.splitlines():
        line = line.rstrip()
        if not line.strip() or _NOISE.search(line):
            continue
        m = _SPEAKER.match(line)
        if m:
            cur = Segment(len(tr.segments), m["spk"].strip(), _to_seconds(m["t"]), "")
            tr.segments.append(cur)
        elif cur is None:
            header.append(line.strip())
        else:
            cur.text = (cur.text + " " + line.strip()).strip()
    if header:
        tr.titulo = header[0]
        _parse_title(header[0], tr)
        for h in header[1:]:
            d = _parse_duration(h)
            if d:
                tr.duracion_seg = d
    tr.segments = [s for s in tr.segments if s.text]
    for i, s in enumerate(tr.segments):
        s.idx = i
    if not tr.segments:
        raise ValueError(
            "No se reconoció el formato: se espera 'Nombre   0:15' seguido del texto."
        )
    return tr


def read_docx(data: bytes) -> Transcript:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        lines.extend(p.text.split("\n"))
    for t in doc.tables:                       # algunas exportaciones usan tablas
        for row in t.rows:
            lines.append("   ".join(c.text.strip() for c in row.cells))
    return parse_text("\n".join(lines))


def read_any(name: str, data: bytes) -> Transcript:
    if name.lower().endswith(".docx"):
        return read_docx(data)
    return parse_text(data.decode("utf-8-sig", errors="replace"))
