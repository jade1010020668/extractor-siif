import io
import socket

import pytest
from docx import Document

from acta_privada import privacy
from acta_privada.docx_builder import build_docx
from acta_privada.llm import OllamaClient
from acta_privada.structure import build_acta
from acta_privada.transcript import read_docx


def test_parser_metadatos(transcript_docx):
    tr = read_docx(transcript_docx)
    assert tr.fecha == (2027, 3, 3)
    assert tr.hora_inicio == (9, 30)
    assert tr.duracion_seg == 15 * 60 + 30
    assert len(tr.segments) == 9            # sin las líneas "ha iniciado/detenido"
    assert "detenido" not in tr.segments[-1].text


def test_modo_reglas(transcript_docx, roster):
    tr = read_docx(transcript_docx)
    acta = build_acta(tr, roster, numero="001")
    assert (acta.hora_inicio, acta.hora_fin) == ("9:30 a.m.", "9:46 a.m.")
    assert "por unanimidad" in acta.apertura
    assert acta.presidente.nombre == "Ana María Pérez Gómez"
    assert [c.nombre for c in acta.comisionados] == ["Carlos Andrés Rojas Lema", "Lucía Fernanda Mora Díaz"]
    assert acta.temas[0].discusion and all(p.startswith("[REVISAR]") for p in acta.temas[0].discusion)


def test_sin_unanimidad_se_avisa(transcript_docx, roster):
    tr = read_docx(transcript_docx)
    tr.segments[3].text = "Tengo dudas."          # Carlos ya no expresa conformidad
    acta = build_acta(tr, roster)
    assert "unanimidad" not in acta.apertura
    assert any("unanimidad" in a for a in acta.avisos)


def test_hablante_desconocido(transcript_docx, roster):
    tr = read_docx(transcript_docx)
    tr.segments[1].speaker = "Persona Desconocida"
    acta = build_acta(tr, roster)
    assert any("no reconocido" in a for a in acta.avisos)


def test_con_ia_y_docx(transcript_docx, roster, fake_ollama):
    url, srv = fake_ollama
    tr = read_docx(transcript_docx)
    llm = OllamaClient(host=url, model="modelo-falso")
    assert llm.has_model()
    acta = build_acta(tr, roster, numero="002", llm=llm)
    tema = acta.temas[0]
    assert tema.punto == 2 and tema.inicio == 4 and tema.fin == 8
    assert "informa que se expidió" in tema.presentacion[0]
    assert tema.decision.startswith("Los Comisionados deciden")
    # el modelo recibe el tratamiento canónico, no el alias de la transcripción
    prompts = " ".join(u for _, u in srv.log)
    assert "La Comisionada Presidente Ana María Pérez Gómez" in prompts

    doc = Document(io.BytesIO(build_docx(acta)))
    assert len(doc.tables) == 5
    texto = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for esperado in ("ACTA No. 002", "Hora de Inicio: 9:30 a.m.", "2.1. Resolución 55 de 2027",
                     "Pretensión:", "Decisión:", "MARTA ELENA SUÁREZ RÍOS"):
        assert esperado in texto or esperado in "\n".join(p.text for p in doc.paragraphs)
    assert doc.core_properties.author == ""


def test_ia_devuelve_basura_cae_a_reglas(transcript_docx, roster, fake_ollama):
    url, srv = fake_ollama
    srv.modo = "basura"
    acta = build_acta(read_docx(transcript_docx), roster, llm=OllamaClient(host=url, model="x"))
    assert any("IA local falló" in a for a in acta.avisos)
    assert acta.temas[0].discusion        # igual se obtiene un borrador


# ───────────────────────────── privacidad ─────────────────────────────
@pytest.mark.parametrize("url", ["http://api.openai.com", "https://example.com:11434",
                                 "http://8.8.8.8:11434", "http://miservidor.empresa.com"])
def test_bloquea_hosts_remotos(url):
    with pytest.raises(privacy.PrivacyError):
        OllamaClient(host=url)


def test_red_local_requiere_autorizacion(monkeypatch):
    with pytest.raises(privacy.PrivacyError):
        OllamaClient(host="http://192.168.1.20:11434")
    monkeypatch.setenv("ACTA_PERMITIR_RED_LOCAL", "1")
    assert OllamaClient(host="http://192.168.1.20:11434").host.startswith("http://192.168")


def test_ignora_proxy_del_sistema(transcript_docx, roster, fake_ollama, monkeypatch):
    url, _ = fake_ollama
    monkeypatch.setenv("HTTP_PROXY", "http://127.0.0.1:9")     # proxy muerto
    monkeypatch.setenv("http_proxy", "http://127.0.0.1:9")
    assert OllamaClient(host=url, model="modelo-falso").has_model()


def test_pipeline_completo_sin_salir_de_loopback(transcript_docx, roster, fake_ollama, monkeypatch):
    """Cualquier conexión a una IP no loopback hace fallar la prueba."""
    url, _ = fake_ollama
    real_connect = socket.socket.connect

    def guard(self, addr, *a, **k):
        host = addr[0] if isinstance(addr, tuple) else str(addr)
        assert host in ("127.0.0.1", "::1", "localhost"), f"conexión externa: {addr}"
        return real_connect(self, addr, *a, **k)

    monkeypatch.setattr(socket.socket, "connect", guard)
    acta = build_acta(read_docx(transcript_docx), roster, llm=OllamaClient(host=url, model="modelo-falso"))
    assert build_docx(acta)
